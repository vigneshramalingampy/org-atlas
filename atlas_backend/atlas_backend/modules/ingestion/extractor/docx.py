from pathlib import Path

from docx import Document as DocxDocument
from loguru import logger

from atlas_backend.modules.ingestion.extractor.base import DocumentExtractor
from atlas_backend.modules.ingestion.extractor.exceptions import (
    DocumentExtractionError,
)
from atlas_backend.modules.ingestion.extractor.model import (
    DocumentMetadata,
    DocumentPage,
    ParsedDocument,
)
from atlas_backend.utils.enums import FileType


class DocxExtractor(DocumentExtractor):
    """Extract text from a .docx Word file.

    Uses python-docx to open the XML-based .docx format and iterate over
    paragraphs. Paragraph boundaries are preserved as double newlines so the
    downstream chunker can recognise logical breaks.

    Tables are extracted row-by-row as pipe-delimited lines.
    """

    async def extract(self, file_path: str) -> ParsedDocument:
        logger.info("Extracting DOCX: {}", file_path)

        path = Path(file_path)
        if not path.exists():
            logger.warning("DOCX file not found: {}", file_path)
            raise DocumentExtractionError(f"File not found: {file_path}")

        try:
            doc = DocxDocument(str(path))
        except Exception as exc:
            logger.error("Failed to open DOCX {}: {}", file_path, exc)
            raise DocumentExtractionError(
                f"Failed to open DOCX {file_path}: {exc}"
            ) from exc

        paragraphs: list[str] = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs.append(text)

        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                line = " | ".join(cells)
                if line.strip():
                    paragraphs.append(line)

        if not paragraphs:
            logger.warning("No extractable text found in DOCX: {}", file_path)
            raise DocumentExtractionError(f"No extractable text found in {file_path}")

        text = "\n\n".join(paragraphs)
        char_count = len(text)
        logger.info(
            "DOCX extracted: {} chars, {} paragraphs from {}",
            char_count,
            len(paragraphs),
            path.name,
        )

        metadata = DocumentMetadata(
            filename=path.name,
            total_pages=1,
            file_type=FileType.DOCX,
        )
        page = DocumentPage(page_number=0, text=text)

        return ParsedDocument(metadata=metadata, pages=[page])
